#!/usr/bin/env python3
"""Empaqueta el Kit de operación (Excel + tablero) para subirlo a Gumroad.

Uso:
    python3 venta/empaquetar_kit.py ruta/a/Kit_de_operacion.xlsx

El Excel NO vive en este repositorio (es público y el Excel es lo que se vende):
se pasa como argumento. El tablero sale de kit-operacion-dashboard/.

Genera dist/kit-de-operacion.zip (dist/ está ignorado por git) con:
  Kit_de_operacion.xlsx
  tablero/index.html, tablero/vendor/*, tablero/LICENSE
  LEEME.txt

Antes de escribir comprueba que el Excel tiene las doce hojas del kit y que el
tablero trae sus dos librerías locales, para que funcione sin internet.
"""
import os
import sys
import zipfile
from datetime import date

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABLERO = os.path.join(RAIZ, "kit-operacion-dashboard")
DIST = os.path.join(RAIZ, "dist")
SALIDA = os.path.join(DIST, "kit-de-operacion.zip")

HOJAS = [
    "Léeme", "01 Diez referencias", "02 Escandallo", "03 Ingeniería de menú",
    "04 Estado de resultados", "05 Los cinco números", "06 Control de barra",
    "07 Tiempo del pase", "08 Ficha del plato", "09 Hora valle",
    "10 Reservas y no-show", "11 Inspector anónimo", "12 Clientes frecuentes",
]
TABLERO_ARCHIVOS = [
    "index.html", "LICENSE",
    "vendor/xlsx.full.min.js", "vendor/chart.umd.js",
    "vendor/LICENSE-xlsx.txt", "vendor/LICENSE-chartjs.md",
]

LEEME = """KIT DE OPERACIÓN
Del libro «Historia de la hostelería y la gastronomía» · Sergio Berriozábal Serrano
Versión {fecha}

QUÉ HAY AQUÍ
  Kit_de_operacion.xlsx   Doce hojas de trabajo. Escribes solo en las celdas azules.
  tablero/index.html      El tablero: tus números de la semana, vistos de un golpe.

CÓMO EMPEZAR
  1. Abre Kit_de_operacion.xlsx. Viene relleno con el restaurante de referencia del
     capítulo 19 para que veas la aritmética funcionando. Sustituye sus cifras por las tuyas.
  2. Guárdalo.
  3. Abre tablero/index.html con doble clic (Chrome, Edge, Safari o Firefox).
  4. Pulsa «Cargar mi Excel» y elige tu archivo.

  El archivo no sale de tu computadora: se lee dentro del navegador. La última carga
  se queda guardada, también sin internet.

EL ORDEN
  Primero las fugas (hojas 01, 02, 07, 10), luego el margen (03, 04, 06, 09), luego el
  cliente (08, 12) y al final el control (05, 11). Una mejora cada tres o cuatro semanas,
  con el número de la anterior ya estabilizado.

QUÉ NO HACE
  No se conecta a tu punto de venta ni a tu contabilidad: los números los escribes tú.
  No es asesoría fiscal ni contable. Las bandas de referencia son las del libro; tu casa
  puede tener las suyas.

LICENCIA
  El Excel y sus textos: uso en tu propio negocio, sin límite de locales propios.
  Prohibida la redistribución, reventa o publicación.
  El código del tablero es MIT (tablero/LICENSE).

AYUDA E INSTALACIÓN EN TU LOCAL
  WhatsApp +52 984 187 8284

Copyright 2026 Sergio Berriozábal Serrano.
"""


def comprobar_excel(ruta):
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("Aviso: sin openpyxl no se comprueban las hojas del Excel (pip install openpyxl).")
        return
    nombres = load_workbook(ruta, read_only=True).sheetnames
    faltan = [h for h in HOJAS if h not in nombres]
    if faltan:
        sys.exit("El Excel no es el del kit; faltan las hojas: " + ", ".join(faltan))


def main(args):
    if len(args) != 1:
        sys.exit(__doc__)
    excel = os.path.abspath(args[0])
    if not os.path.isfile(excel):
        sys.exit("No existe el Excel: " + excel)
    comprobar_excel(excel)
    faltan = [a for a in TABLERO_ARCHIVOS if not os.path.isfile(os.path.join(TABLERO, a))]
    if faltan:
        sys.exit("Faltan archivos del tablero: " + ", ".join(faltan))

    os.makedirs(DIST, exist_ok=True)
    with zipfile.ZipFile(SALIDA, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(excel, "Kit_de_operacion.xlsx")
        for a in TABLERO_ARCHIVOS:
            z.write(os.path.join(TABLERO, a), "tablero/" + a)
        z.writestr("LEEME.txt", LEEME.format(fecha=date.today().isoformat()))
    print(f"{SALIDA} · {os.path.getsize(SALIDA) // 1024} KB")


if __name__ == "__main__":
    main(sys.argv[1:])
