#!/usr/bin/env python3
"""
censo_extraccion.py — FASE 0 (censo) y FASE 1 (extracción total) del
Compilador de Contexto Universal.

Uso:
  python3 censo_extraccion.py --censo   <carpeta_fuente> --salida <trabajo>
  python3 censo_extraccion.py --extraer <carpeta_fuente> --salida <trabajo> [--incluir lista.txt]

--censo   escribe <trabajo>/CENSO.md y <trabajo>/censo.json:
          archivos por tipo y subcarpeta, fechas, duplicados exactos (SHA-256),
          familias de versiones (v1/v2/final/FINAL2/copia de) y propuesta de
          qué entra y qué se descarta.
--extraer lee cada archivo aprobado con la herramienta de su formato y vuelca
          el texto a <trabajo>/extracciones/<ruta_plana>.md. Lo ilegible va a
          <trabajo>/ILEGIBLES.md. Nada más: el script no interpreta.

La carpeta fuente nunca se escribe. Solo se lee.
"""
import argparse, hashlib, json, os, re, subprocess, sys
from datetime import datetime
from pathlib import Path

TEXTO = {".md", ".txt", ".csv", ".json", ".html", ".htm", ".rtf"}
DOCX, PDF, PPTX, XLSX = {".docx"}, {".pdf"}, {".pptx"}, {".xlsx", ".xlsm"}
IMG = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}
IGNORAR_DIR = {"CONTEXTO", ".git", "node_modules", "__pycache__", ".DS_Store"}
VERSION_RE = re.compile(
    r"(?i)[\s_\-\(]*(v\s?\d+(\.\d+)*|final\s*\d*|definitiv[oa]\s*\d*|copia(\s+de)?|copy|"
    r"borrador|draft|old|antigu[oa]|viej[oa]|\(\d+\)|nuevo|new|rev\s?\d*)[\s_\-\)]*")


def sha256(p, bloque=1 << 20):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(bloque), b""):
            h.update(b)
    return h.hexdigest()


def familia(nombre):
    """Quita marcas de versión para agrupar 'Precios v2.docx' con 'Precios FINAL.docx'."""
    base = Path(nombre).stem
    base = VERSION_RE.sub(" ", base)
    base = re.sub(r"[\s_\-]+", " ", base).strip().lower()
    return base or Path(nombre).stem.lower()


def listar(fuente):
    fuente = Path(fuente)
    out = []
    for raiz, dirs, files in os.walk(fuente):
        dirs[:] = [d for d in dirs if d not in IGNORAR_DIR and not d.startswith(".")]
        for f in files:
            if f.startswith(".") or f.startswith("~$"):
                continue
            p = Path(raiz) / f
            try:
                st = p.stat()
            except OSError:
                continue
            out.append({
                "ruta": str(p.relative_to(fuente)),
                "carpeta": str(p.parent.relative_to(fuente)) or ".",
                "nombre": f,
                "ext": p.suffix.lower(),
                "bytes": st.st_size,
                "modificado": datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d"),
                "sha256": sha256(p) if st.st_size else "vacio",
                "familia": familia(f),
            })
    return sorted(out, key=lambda x: x["ruta"])


def censo(fuente, salida):
    salida = Path(salida); salida.mkdir(parents=True, exist_ok=True)
    arch = listar(fuente)
    por_hash, por_familia, por_tipo, por_carpeta = {}, {}, {}, {}
    for a in arch:
        por_hash.setdefault(a["sha256"], []).append(a)
        por_familia.setdefault((a["carpeta"], a["familia"], a["ext"]), []).append(a)
        por_tipo[a["ext"] or "(sin ext)"] = por_tipo.get(a["ext"] or "(sin ext)", 0) + 1
        por_carpeta[a["carpeta"]] = por_carpeta.get(a["carpeta"], 0) + 1

    duplicados = [g for h, g in por_hash.items() if len(g) > 1 and h != "vacio"]
    familias = [g for g in por_familia.values() if len(g) > 1]

    # Propuesta: entra la copia más reciente de cada grupo; el resto, historial.
    descartes = set()
    for g in duplicados:
        g_sorted = sorted(g, key=lambda x: x["modificado"], reverse=True)
        for d in g_sorted[1:]:
            descartes.add((d["ruta"], f"duplicado exacto de {g_sorted[0]['ruta']}"))
    for g in familias:
        g_sorted = sorted(g, key=lambda x: x["modificado"], reverse=True)
        for d in g_sorted[1:]:
            if not any(d["ruta"] == r for r, _ in descartes):
                descartes.add((d["ruta"], f"versión anterior de la familia «{g_sorted[0]['nombre']}» — pasa a historial"))
    legibles = {*TEXTO, *DOCX, *PDF, *PPTX, *XLSX, *IMG}
    for a in arch:
        if a["ext"] not in legibles:
            descartes.add((a["ruta"], f"formato no extraíble ({a['ext'] or 'sin extensión'})"))
        if a["sha256"] == "vacio":
            descartes.add((a["ruta"], "archivo de 0 bytes"))
    desc_map = {r: m for r, m in descartes}
    entran = [a["ruta"] for a in arch if a["ruta"] not in desc_map]

    L = [f"# CENSO DE LA CARPETA — {fuente}", "",
         f"Compilado: {datetime.now():%Y-%m-%d %H:%M} · Archivos: {len(arch)} · "
         f"Entran: {len(entran)} · Descartes propuestos: {len(desc_map)}", "",
         "## Archivos por tipo", "", "| Tipo | Nº |", "|---|---|"]
    L += [f"| {t} | {n} |" for t, n in sorted(por_tipo.items(), key=lambda x: -x[1])]
    L += ["", "## Archivos por subcarpeta", "", "| Carpeta | Nº |", "|---|---|"]
    L += [f"| {c} | {n} |" for c, n in sorted(por_carpeta.items())]
    L += ["", "## Inventario completo", "", "| Ruta | Modificado | Tamaño | Entra |", "|---|---|---|---|"]
    L += [f"| {a['ruta']} | {a['modificado']} | {a['bytes']:,} | "
          f"{'✅' if a['ruta'] in entran else '❌ ' + desc_map[a['ruta']]} |" for a in arch]
    L += ["", f"## Duplicados exactos (SHA-256) — {len(duplicados)} grupos", ""]
    for g in duplicados:
        L.append("- " + " ≡ ".join(x["ruta"] for x in g))
    L += ["", f"## Familias de versiones — {len(familias)} grupos", ""]
    for g in familias:
        g_sorted = sorted(g, key=lambda x: x["modificado"], reverse=True)
        L.append(f"- **{g_sorted[0]['familia']}** → vigente propuesta: `{g_sorted[0]['ruta']}` "
                 f"({g_sorted[0]['modificado']}); historial: " +
                 ", ".join(f"`{x['ruta']}` ({x['modificado']})" for x in g_sorted[1:]))
    L += ["", "## Propuesta", "",
          "La vigente de cada familia es la más reciente por fecha de modificación. "
          "Si una versión anterior está marcada como APROBADA dentro del texto, mándala "
          "de vuelta: el censo lee metadatos, no criterio. **Espera aprobación antes de extraer**, "
          "salvo que el usuario haya dicho «hazlo directo»."]
    (salida / "CENSO.md").write_text("\n".join(L), encoding="utf-8")
    (salida / "censo.json").write_text(json.dumps(
        {"fuente": str(fuente), "archivos": arch, "entran": entran,
         "descartes": [{"ruta": r, "motivo": m} for r, m in sorted(descartes)]},
        ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"CENSO: {len(arch)} archivos · {len(entran)} entran · {len(desc_map)} descartes → {salida/'CENSO.md'}")


# ---------------- extracción ----------------
def ex_docx(p):
    import docx
    d = docx.Document(p); out = []
    for para in d.paragraphs:
        if para.text.strip():
            st = para.style.name.lower() if para.style is not None else ""
            pref = "## " if st.startswith("heading 1") or st.startswith("título 1") else \
                   "### " if st.startswith("heading") or st.startswith("título") else ""
            out.append(pref + para.text.strip())
    for t in d.tables:
        out.append("")
        for i, row in enumerate(t.rows):
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            out.append("| " + " | ".join(cells) + " |")
            if i == 0:
                out.append("|" + "---|" * len(cells))
    return "\n".join(out)


def ex_pdf(p):
    r = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True)
    txt = r.stdout if r.returncode == 0 else ""
    if len(txt.strip()) > 200:
        return txt
    # Escaneado: OCR página a página
    try:
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            subprocess.run(["pdftoppm", "-r", "200", "-png", str(p), f"{td}/pg"], check=True, capture_output=True)
            paginas = sorted(Path(td).glob("pg*.png"))
            out = []
            for pg in paginas:
                o = subprocess.run(["tesseract", str(pg), "-", "-l", "spa+eng"], capture_output=True, text=True)
                out.append(o.stdout)
            return ("[OCR]\n" + "\n\n".join(out)) if any(x.strip() for x in out) else txt
    except Exception as e:
        return txt + f"\n[OCR FALLIDO: {e}]"


def ex_pptx(p):
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(p).slides, 1):
        out.append(f"\n## Diapositiva {i}")
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for j, row in enumerate(sh.table.rows):
                    cells = [c.text.strip() for c in row.cells]
                    out.append("| " + " | ".join(cells) + " |")
                    if j == 0:
                        out.append("|" + "---|" * len(cells))
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            out.append("_Notas:_ " + s.notes_slide.notes_text_frame.text.strip())
    return "\n".join(out)


def ex_xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n## Hoja: {ws.title}")
        filas = 0
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None for v in row):
                continue
            out.append("| " + " | ".join("" if v is None else str(v) for v in row) + " |")
            filas += 1
            if filas == 1:
                out.append("|" + "---|" * len(row))
            if filas > 400:
                out.append("| … (truncado a 400 filas; revisar la fuente) |"); break
    return "\n".join(out)


def ex_img(p):
    o = subprocess.run(["tesseract", str(p), "-", "-l", "spa+eng"], capture_output=True, text=True)
    return ("[OCR]\n" + o.stdout) if o.stdout.strip() else ""


def extraer(fuente, salida, incluir=None):
    fuente, salida = Path(fuente), Path(salida)
    cj = json.loads((salida / "censo.json").read_text(encoding="utf-8"))
    rutas = cj["entran"]
    if incluir:
        rutas = [l.strip() for l in Path(incluir).read_text(encoding="utf-8").splitlines() if l.strip()]
    dest = salida / "extracciones"; dest.mkdir(parents=True, exist_ok=True)
    ilegibles, manifest = [], []
    for r in rutas:
        p = fuente / r; ext = p.suffix.lower()
        try:
            if ext in TEXTO: txt = p.read_text(encoding="utf-8", errors="replace")
            elif ext in DOCX: txt = ex_docx(p)
            elif ext in PDF: txt = ex_pdf(p)
            elif ext in PPTX: txt = ex_pptx(p)
            elif ext in XLSX: txt = ex_xlsx(p)
            elif ext in IMG: txt = ex_img(p)
            else: raise ValueError("formato no soportado")
            if len(txt.strip()) < 20:
                raise ValueError("sin texto legible (¿escaneo sin OCR, imagen sin texto?)")
        except Exception as e:
            ilegibles.append((r, str(e))); continue
        plano = re.sub(r"[^\w\-.]+", "_", r)
        cab = (f"---\nfuente: {r}\nmodificado: "
               f"{datetime.fromtimestamp(p.stat().st_mtime):%Y-%m-%d}\n---\n\n")
        (dest / f"{plano}.md").write_text(cab + txt, encoding="utf-8")
        manifest.append((r, f"extracciones/{plano}.md", len(txt)))
    (salida / "ILEGIBLES.md").write_text(
        "# Archivos ilegibles\n\n| Archivo | Motivo |\n|---|---|\n" +
        "\n".join(f"| {r} | {m} |" for r, m in ilegibles) + ("\n\n(ninguno)" if not ilegibles else ""),
        encoding="utf-8")
    (salida / "MANIFEST_EXTRACCION.md").write_text(
        "# Manifest de extracción\n\n| Archivo fuente | Extracción | Caracteres |\n|---|---|---|\n" +
        "\n".join(f"| {r} | {e} | {n:,} |" for r, e, n in manifest), encoding="utf-8")
    print(f"EXTRACCIÓN: {len(manifest)} leídos · {len(ilegibles)} ilegibles → {dest}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--censo", metavar="FUENTE")
    g.add_argument("--extraer", metavar="FUENTE")
    ap.add_argument("--salida", required=True)
    ap.add_argument("--incluir", help="lista de rutas aprobadas (una por línea), opcional")
    a = ap.parse_args()
    if a.censo: censo(a.censo, a.salida)
    else: extraer(a.extraer, a.salida, a.incluir)
